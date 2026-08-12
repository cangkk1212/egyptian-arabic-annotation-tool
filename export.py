#!/usr/bin/env python3
"""导出标注统计到 Excel — 按用户汇总已标注/跳过的音频"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

SCRIPT_DIR = Path(__file__).parent.resolve()
CONFIG_PATH = SCRIPT_DIR / "config.json"

# 读取配置获取标注目录
config = {}
if CONFIG_PATH.exists():
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        config = json.load(f)
ANNOTATIONS_DIR = config.get("annotations_dir", "/home/ck/annotations")

OUTPUT = SCRIPT_DIR / "annotation_export.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Annotations"

# Header
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color="1e1e2e", end_color="1e1e2e", fill_type="solid")
header_font_white = Font(bold=True, size=11, color="ffffff")
headers = ["User", "Folder", "Audio", "Duration (s)", "Status"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# Scan annotations
ann_dir = Path(ANNOTATIONS_DIR)
row = 2
count = 0
for f in sorted(ann_dir.rglob("*.json")):
    if f.name in ("active_users.json", "assignments.json"):
        continue
    try:
        d = json.load(open(f, "r", encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        continue
    status = d.get("status", "pending")
    if status not in ("annotated", "skipped"):
        continue
    user = d.get("annotated_by") or d.get("skipped_by") or ""
    ws.cell(row=row, column=1, value=user)
    ws.cell(row=row, column=2, value=d.get("folder", ""))
    ws.cell(row=row, column=3, value=d.get("audio", f.stem))
    ws.cell(row=row, column=4, value=round(d.get("duration", 0), 1))
    ws.cell(row=row, column=5, value=status)
    row += 1
    count += 1

# Auto-adjust column widths
for col_letter in ["A", "B", "C", "D", "E"]:
    max_len = 0
    for cell in ws[col_letter]:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

# Freeze header
ws.freeze_panes = "A2"

wb.save(str(OUTPUT))
print(f"✅ Exported {count} rows to {OUTPUT}")
